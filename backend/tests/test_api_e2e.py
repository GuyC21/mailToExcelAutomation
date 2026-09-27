"""End-to-end API flow in mock mode: email in -> DB -> Excel -> preview."""
import base64
import os
from email.message import EmailMessage

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from config import get_settings
from main import app
from repositories.excel_layout import DOCUMENTS_SHEET, LINES_SHEET
from repositories.excel_repository import ExcelLockedError, ExcelRepository

PDF = b"%PDF-1.4\n% test\n"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client


def test_health_reports_mock_provider(client):
    assert client.get("/api/health").json()["extraction_providers"] == ["mock"]


def test_multipart_email_processes_pdf_and_skips_other_attachments(client):
    response = client.post(
        "/api/email/inbound",
        data={"from": "billing@medipharm-care.co.il", "to": "invoices@goldencare.co.il",
              "subject": "טופס התחשבנות אוגוסט", "text": "מצורף הטופס", "headers": "X-Mailer: Outlook"},
        files=[("attachments", ("טופס.pdf", PDF, "application/pdf")),
               ("attachments", ("notes.txt", b"hello", "text/plain"))],
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["email"]["subject"] == "טופס התחשבנות אוגוסט"
    assert body["email"]["headers"]["X-Mailer"] == "Outlook"
    assert body["email"]["message_id"]
    statuses = [(r["filename"], r["status"]) for r in body["results"]]
    assert statuses == [("טופס.pdf", "VALID"), ("notes.txt", "SKIPPED")]
    assert body["results"][0]["excel"]["written"] is True
    assert body["results"][0]["provider"] == "mock"


def test_eml_and_json_formats(client):
    message = EmailMessage()
    message["From"], message["To"], message["Subject"] = "a@b.co.il", "invoices@goldencare.co.il", "חשבונית"
    message.set_content("גוף")
    message.add_attachment(PDF, maintype="application", subtype="pdf", filename="form.pdf")
    eml = client.post("/api/email/inbound/eml", files={"file": ("m.eml", message.as_bytes(), "message/rfc822")})
    assert eml.status_code == 200 and eml.json()["results"][0]["status"] == "VALID"

    payload = {"from": "a@b.co.il", "subject": "json", "attachments": [
        {"filename": "form.pdf", "content_base64": base64.b64encode(PDF).decode()}]}
    assert client.post("/api/email/inbound/json", json=payload).json()["results"][0]["status"] == "VALID"


def test_sandbox_flags_math_errors_and_rejects_non_documents(client):
    flagged = client.post("/api/ingestion/upload", files={"file": ("טופס תקול.pdf", PDF, "application/pdf")}).json()
    assert flagged["status"] == "NEEDS_REVIEW"
    assert {i["code"] for i in flagged["issues"] if i["severity"] == "error"} == {"LINE_MATH_MISMATCH", "SUBTOTAL_MISMATCH"}
    fake = client.post("/api/ingestion/upload", files={"file": ("x.pdf", b"not a pdf", "application/pdf")})
    assert fake.status_code == 415


def test_workbook_accumulates_documents_and_lines(client):
    workbook = load_workbook(get_settings().excel_path)
    documents = list(workbook[DOCUMENTS_SHEET].iter_rows(min_row=2, values_only=True))
    assert len(documents) == 4
    assert workbook[DOCUMENTS_SHEET].sheet_view.rightToLeft is True
    assert workbook[LINES_SHEET].max_row - 1 == 8  # 2 mock lines x 4 documents
    preview = client.get("/api/excel/preview?limit=2").json()
    assert preview["total"] == 4 and len(preview["rows"]) == 2
    assert preview["counts"] == {"VALID": 3, "NEEDS_REVIEW": 1, "EXTRACTION_FAILED": 0}
    assert client.get("/api/excel/download").status_code == 200


def test_locked_workbook_keeps_rows_pending_then_syncs(client, monkeypatch):
    def locked(self, workbook):
        raise ExcelLockedError("locked")
    monkeypatch.setattr(ExcelRepository, "_save", locked)
    result = client.post("/api/ingestion/upload", files={"file": ("a.pdf", PDF, "application/pdf")}).json()
    assert result["excel"] == {"written": False, "message": "locked"}
    monkeypatch.undo()
    assert client.post("/api/excel/sync").json()["synced"] == [result["ingestion_id"]]
    assert client.get("/api/excel/preview").json()["total"] == 5


def test_technical_prompt_shows_real_envelope(client):
    prompt = client.get("/api/prompts/active").json()
    technical = client.get(f"/api/prompts/{prompt['id']}/technical").json()["technical_prompt"]
    assert prompt["content"] in technical and '"vat_amount"' in technical


def test_legacy_workbook_is_archived(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "book.xlsx"
    Workbook().save(path)
    ExcelRepository(str(path)).ensure_workbook()
    assert DOCUMENTS_SHEET in load_workbook(path).sheetnames
    assert any(name.startswith("book.legacy-") for name in os.listdir(tmp_path))
