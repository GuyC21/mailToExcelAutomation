"""Regression tests for the Sprint 3 bug review (findings B01-B17).

Each test reproduces the trigger described in the review and asserts the
acceptance check, so a fix cannot silently regress. Runs offline (mock
provider / stub providers), on the shared test database from ``conftest``.
"""
import asyncio
import base64
import copy
import json
import os
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from starlette.applications import Starlette
from starlette.responses import PlainTextResponse
from starlette.routing import Route

import services.regression.labeling as labeling
from api.uploads import BodySizeLimitMiddleware
from config import get_settings
from database import async_session
from main import app
from models.ingestion import DocumentIngestion
from repositories.excel_layout import DOCUMENTS_SHEET, DOCUMENT_COLUMNS, LINES_SHEET
from schemas.email_payload import JsonEmailRequest
from schemas.extraction import DocumentExtraction
from services.email_parser import parse_json
from services.excel_sync import get_excel_repository, sync_pending
from services.extraction.extractor import DocumentExtractor
from services.extraction.providers.base import ExtractionProvider, ProviderError, ProviderResponse
from services.extraction.providers.gemini_provider import GeminiProvider
from services.ingestion_pipeline import ingest_email
from services.regression.runner import RegressionRunner
from services.regression.scorer import score_extraction
from tests.fixtures import sample_documents as samples

PDF = b"%PDF-1.4\n% test\n"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client


def _json_email(message_id=None, pdf=PDF, subject="חשבונית", filename="form.pdf"):
    payload = {"from": "billing@supplier.co.il", "subject": subject,
               "attachments": [{"filename": filename, "content_base64": base64.b64encode(pdf).decode()}]}
    if message_id:
        payload["message_id"] = message_id
    return payload


async def _count_documents() -> int:
    async with async_session() as session:
        return (await session.execute(select(func.count()).select_from(DocumentIngestion))).scalar()


def _workbook_ids():
    workbook = load_workbook(get_settings().excel_path)
    try:
        return sorted(row[0] for row in workbook[DOCUMENTS_SHEET].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# B03 - email retries must not duplicate invoice records
# ---------------------------------------------------------------------------

def test_b03_replayed_email_creates_exactly_one_record(client):
    message_id = f"<{uuid.uuid4().hex}@supplier.co.il>"
    before = asyncio.run(_count_documents())
    first = client.post("/api/email/inbound/json", json=_json_email(message_id)).json()["results"][0]
    second = client.post("/api/email/inbound/json", json=_json_email(message_id)).json()["results"][0]

    assert first["duplicate"] is False and first["status"] == "VALID"
    assert second["duplicate"] is True
    assert second["ingestion_id"] == first["ingestion_id"]
    assert asyncio.run(_count_documents()) == before + 1
    assert _workbook_ids().count(first["ingestion_id"]) == 1


def test_b03_new_attachment_in_the_same_email_still_processes(client):
    message_id = f"<{uuid.uuid4().hex}@supplier.co.il>"
    client.post("/api/email/inbound/json", json=_json_email(message_id))
    other = PDF + b"% a different document\n"
    result = client.post("/api/email/inbound/json", json=_json_email(message_id, pdf=other)).json()["results"][0]
    assert result["duplicate"] is False and result["ingestion_id"]


def test_b03_replay_without_message_id_is_detected_from_content(client):
    subject = f"no-id {uuid.uuid4().hex}"
    first = client.post("/api/email/inbound/json", json=_json_email(subject=subject)).json()["results"][0]
    second = client.post("/api/email/inbound/json", json=_json_email(subject=subject)).json()["results"][0]
    assert second["duplicate"] is True and second["ingestion_id"] == first["ingestion_id"]


def test_b03_concurrent_retries_create_one_record(client):
    message_id = f"<{uuid.uuid4().hex}@supplier.co.il>"
    payload = parse_json(JsonEmailRequest(**_json_email(message_id)))

    async def deliver():
        async with async_session() as session:
            _email, results = await ingest_email(session, payload.model_copy(deep=True))
            return results[0]

    async def both():
        return await asyncio.gather(deliver(), deliver())

    before = asyncio.run(_count_documents())
    results = asyncio.run(both())
    assert asyncio.run(_count_documents()) == before + 1
    assert sorted(r.duplicate for r in results) == [False, True]


# ---------------------------------------------------------------------------
# B04 / B05 - the workbook is rebuilt from the database, never merged blindly
# ---------------------------------------------------------------------------

def _write_foreign_workbook(path, ids):
    workbook = Workbook()
    workbook.active.title = DOCUMENTS_SHEET
    workbook.active.append([c.header for c in DOCUMENT_COLUMNS])
    for ingestion_id in ids:
        workbook.active.append([ingestion_id, None, "תקין", None, None, None, None, "old.pdf", "ספק ישן"])
    workbook.create_sheet(LINES_SHEET)
    workbook.save(path)


def test_b04_foreign_workbook_is_archived_and_rebuilt(client):
    client.post("/api/ingestion/upload", files={"file": ("a.pdf", PDF, "application/pdf")})
    path = get_settings().excel_path
    db_count = asyncio.run(_count_documents())
    _write_foreign_workbook(path, range(1, db_count + 30))  # overlapping, unrelated ids

    async def run():
        async with async_session() as session:
            return await sync_pending(session)

    result = asyncio.run(run())
    assert result["error"] is None
    assert len(_workbook_ids()) == db_count  # exactly the database's records
    workbook = load_workbook(path)
    assert "ספק ישן" not in [row[8] for row in workbook[DOCUMENTS_SHEET].iter_rows(min_row=2, values_only=True)]
    folder = os.path.dirname(path)
    assert any(".orphan-" in name for name in os.listdir(folder))


def test_b05_deleted_workbook_is_restored_with_full_history(client):
    db_count = asyncio.run(_count_documents())
    assert db_count > 0
    os.remove(get_settings().excel_path)

    async def run():
        async with async_session() as session:
            return await sync_pending(session)

    asyncio.run(run())
    assert len(_workbook_ids()) == db_count
    assert get_excel_repository().read_identity()


def test_b05_archived_layout_is_rebuilt_with_history(tmp_path, client):
    db_count = asyncio.run(_count_documents())
    Workbook().save(get_settings().excel_path)  # outdated layout -> archived as legacy

    async def run():
        async with async_session() as session:
            return await sync_pending(session)

    asyncio.run(run())
    assert len(_workbook_ids()) == db_count


# ---------------------------------------------------------------------------
# B07 - manual sync responses serialize
# ---------------------------------------------------------------------------

def test_b07_sync_response_serializes_when_nothing_is_pending(client):
    response = client.post("/api/excel/sync")
    assert response.status_code == 200
    assert response.json() == {"synced": [], "pending": 0, "error": None}


# ---------------------------------------------------------------------------
# B09 - external text is never written as a formula
# ---------------------------------------------------------------------------

def test_b09_formula_shaped_subject_is_stored_as_text(client):
    subject = f"=1+1 {uuid.uuid4().hex[:6]}"
    client.post("/api/email/inbound/json", json=_json_email(f"<{uuid.uuid4().hex}@x>", subject=subject))
    workbook = load_workbook(get_settings().excel_path)
    subject_column = [c.header for c in DOCUMENT_COLUMNS].index("נושא המייל") + 1
    cells = [row[subject_column - 1] for row in workbook[DOCUMENTS_SHEET].iter_rows(min_row=2)]
    match = next(cell for cell in cells if cell.value == subject)
    assert match.data_type == "s"


# ---------------------------------------------------------------------------
# B08 / B14 - dashboard scope and currencies
# ---------------------------------------------------------------------------

async def _add_record(**kwargs):
    async with async_session() as session:
        record = DocumentIngestion(channel="email", provider="gemini", model="g", status="VALID",
                                   excel_synced=True, **kwargs)
        session.add(record)
        await session.commit()


def test_b08_mock_and_sandbox_records_are_excluded_by_default(client):
    operational = client.get("/api/dashboard/stats").json()
    everything = client.get("/api/dashboard/stats?scope=all").json()
    assert operational["scope"] == "operational"
    assert operational["excluded_test_documents"] > 0
    assert everything["total_documents"] == operational["total_documents"] + operational["excluded_test_documents"]
    assert client.get("/api/dashboard/stats?scope=bogus").status_code == 422


def test_b14_currencies_are_never_added_together(client):
    tag = uuid.uuid4().hex[:6]
    asyncio.run(_add_record(extracted_data={"supplier_name": f"USD {tag}", "total_amount": 100.0, "currency": "USD"}))
    asyncio.run(_add_record(extracted_data={"supplier_name": f"ILS {tag}", "total_amount": 100.0, "currency": "₪"}))
    asyncio.run(_add_record(extracted_data={"supplier_name": f"None {tag}", "total_amount": None, "currency": "ILS"}))
    stats = client.get("/api/dashboard/stats").json()
    totals = {t["currency"]: t["amount"] for t in stats["totals_by_currency"]}
    assert totals["USD"] == 100.0
    assert totals["ILS"] >= 100.0 and "₪" not in totals
    assert stats["top_suppliers_currency"] in totals
    assert "total_amount_processed" not in stats  # the old mixed-currency sum is gone


# ---------------------------------------------------------------------------
# B10 / B16 - number normalisation
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw, expected", [
    ("1.234,56", 1234.56), ("1234,56", 1234.56), ("1 234,56", 1234.56), ("6,300.00 ₪", 6300.0),
    ("1,234", 1234.0), ("-500", -500.0), ("12,5", 12.5),
])
def test_b10_supported_formats_keep_their_magnitude(raw, expected):
    assert DocumentExtraction.model_validate({"total_amount": raw}).total_amount == expected


def test_b10_decimal_comma_is_noted_and_ambiguous_grouping_rejected():
    doc = DocumentExtraction.model_validate({"subtotal": "1.234,56", "total_amount": "1,23,456"})
    assert doc.total_amount is None
    assert any("פסיק עשרוני" in note for note in doc.extraction_notes)
    assert any("1,23,456" in note for note in doc.extraction_notes)


@pytest.mark.parametrize("raw", ["NaN", "Infinity", "-Infinity", "1e400", float("nan"), float("inf")])
def test_b16_nonfinite_amounts_are_never_valid(raw):
    from services.validation.business_rules import derive_status, validate_document

    data = copy.deepcopy(samples.SHL_VALID)
    data["total_amount"] = raw
    doc = DocumentExtraction.model_validate(data)
    assert doc.total_amount is None
    assert derive_status(validate_document(doc)) == "NEEDS_REVIEW"
    json.dumps(doc.model_dump(), allow_nan=False)  # serializable


# ---------------------------------------------------------------------------
# B11 - bounded uploads
# ---------------------------------------------------------------------------

def _tiny_app(max_bytes):
    async def echo(request):
        body = await request.body()
        return PlainTextResponse(str(len(body)))
    small = Starlette(routes=[Route("/", echo, methods=["POST"])])
    small.add_middleware(BodySizeLimitMiddleware, max_bytes=max_bytes)
    return small


def test_b11_request_body_limit_by_header_and_by_stream():
    with TestClient(_tiny_app(100)) as small:
        assert small.post("/", content=b"x" * 50).text == "50"
        assert small.post("/", content=b"x" * 500).status_code == 413

        def chunks():
            for _ in range(10):
                yield b"x" * 50
        assert small.post("/", content=chunks()).status_code == 413  # chunked, no Content-Length


def test_b11_sandbox_and_labeling_uploads_are_capped(client, monkeypatch, tmp_path):
    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    monkeypatch.setattr(labeling, "suite_dir", lambda: tmp_path)
    sandbox = client.post("/api/ingestion/upload", files={"file": ("a.pdf", PDF, "application/pdf")})
    assert sandbox.status_code == 413
    label = client.post("/api/regression/cases",
                        data={"case_name": "big", "expected_json": json.dumps(samples.SHL_VALID)},
                        files={"file": ("big.pdf", PDF, "application/pdf")})
    assert label.status_code == 413
    assert not list(tmp_path.iterdir())


def test_b11_oversized_email_attachment_is_skipped_not_buffered(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    multipart = client.post("/api/email/inbound", data={"from": "a@b.co.il", "subject": "big"},
                            files=[("attachments", ("big.pdf", PDF, "application/pdf"))])
    assert multipart.status_code == 200
    assert multipart.json()["results"][0]["status"] == "SKIPPED"
    as_json = client.post("/api/email/inbound/json", json=_json_email(f"<{uuid.uuid4().hex}@x>"))
    assert as_json.json()["results"][0]["status"] == "SKIPPED"


def test_b11_too_many_attachments_is_rejected(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "max_email_attachments", 1)
    response = client.post("/api/email/inbound", data={"from": "a@b.co.il"},
                           files=[("attachments", ("a.pdf", PDF, "application/pdf")),
                                  ("attachments", ("b.pdf", PDF, "application/pdf"))])
    assert response.status_code == 413


# ---------------------------------------------------------------------------
# B12 - optional Backoffice key
# ---------------------------------------------------------------------------

def test_b12_backoffice_key_protects_reads_and_writes(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "backoffice_api_key", "s3cret")
    assert client.get("/api/dashboard/stats").status_code == 401
    assert client.get("/api/prompts/").status_code == 401
    assert client.get("/api/excel/download").status_code == 401
    assert client.delete("/api/regression/cases/anything").status_code == 401
    assert client.get("/api/dashboard/stats", headers={"X-API-Key": "wrong"}).status_code == 401
    assert client.get("/api/dashboard/stats", headers={"X-API-Key": "s3cret"}).status_code == 200
    assert client.get("/api/excel/download?api_key=s3cret").status_code == 200  # plain links
    assert client.get("/api/health").status_code == 200
    # Webhook: open only with no secrets; the Backoffice key is accepted too.
    assert client.post("/api/email/inbound/json", json=_json_email()).status_code == 401
    ok = client.post("/api/email/inbound/json", json=_json_email(f"<{uuid.uuid4().hex}@x>"),
                     headers={"X-API-Key": "s3cret"})
    assert ok.status_code == 200


def test_b12_no_key_configured_keeps_local_demo_open(client):
    assert get_settings().backoffice_api_key == ""
    assert client.get("/api/dashboard/stats").status_code == 200


# ---------------------------------------------------------------------------
# B17 - saved document preview is inline
# ---------------------------------------------------------------------------

def test_b17_case_document_is_served_inline(client, monkeypatch, tmp_path):
    monkeypatch.setattr(labeling, "suite_dir", lambda: tmp_path)
    client.post("/api/regression/cases", data={"case_name": "review_case", "expected_json": json.dumps(samples.SHL_VALID)},
                files={"file": ("review_case.pdf", PDF, "application/pdf")})
    response = client.get("/api/regression/cases/review_case/document")
    assert response.status_code == 200
    assert response.headers["content-disposition"].startswith("inline")
    assert response.headers["content-type"].startswith("application/pdf")


# ---------------------------------------------------------------------------
# B06 - wrong critical fields fail regardless of the aggregate score
# ---------------------------------------------------------------------------

def _normalised(data):
    return DocumentExtraction.model_validate(data).model_dump()


@pytest.mark.parametrize("field, wrong", [
    ("total_amount", 0), ("supplier_tax_id", "512998411"), ("document_number", "SLH-88422"), ("currency", "USD"),
])
def test_b06_critical_field_mismatch_is_reported(field, wrong):
    expected = _normalised(samples.SHL_VALID)
    actual = dict(expected, **{field: wrong})
    report = score_extraction(expected, actual)
    assert field in report.critical_failures


def test_b06_zero_total_fails_the_case_even_with_a_high_score(tmp_path):
    expected = _normalised(samples.SHL_VALID)
    (tmp_path / "shl.pdf").write_bytes(PDF)
    (tmp_path / "shl_expected.json").write_text(json.dumps(expected, ensure_ascii=False), encoding="utf-8")
    runner = RegressionRunner(DocumentExtractor([_Stub(dict(expected, total_amount=0))]), pass_threshold=80.0)
    report = asyncio.run(runner.run(tmp_path))
    result = report.results[0]
    assert result.score.overall_score >= 80.0  # the aggregate alone would have passed
    assert result.passed is False
    assert "total_amount" in result.score.critical_failures


def test_b06_money_tolerance_is_absolute_only():
    expected = _normalised(samples.SHL_VALID)
    actual = copy.deepcopy(expected)
    actual["line_items"][0]["line_total"] = 6250  # printed 6,300
    report = score_extraction(expected, actual)
    line = next(c for c in report.line_item_comparisons if c.status == "matched" and c.expected_index == 0)
    assert not next(f for f in line.fields if f.field == "line_total").matched
    rounding = dict(expected, total_amount=expected["total_amount"] + 0.4)
    assert score_extraction(expected, rounding).critical_failures == []


# ---------------------------------------------------------------------------
# B15 - failed extractions count against headline accuracy
# ---------------------------------------------------------------------------

class _Stub(ExtractionProvider):
    name = "stub"

    def __init__(self, payload=None, fail_for=None):
        self._payload, self._fail_for = payload, fail_for or set()

    def is_configured(self):
        return True

    def extract(self, system_prompt, file_path, mime_type):
        if os.path.basename(file_path).split(".")[0] in self._fail_for:
            raise ProviderError("boom", attempts=[{"provider": "stub", "model": "s", "error": "boom"}])
        return ProviderResponse(raw_text=json.dumps(self._payload), model="stub-1")


def test_b15_failure_is_not_hidden_by_average(tmp_path):
    expected = _normalised(samples.SHL_VALID)
    for stem in ("good", "bad"):
        (tmp_path / f"{stem}.pdf").write_bytes(PDF)
        (tmp_path / f"{stem}_expected.json").write_text(json.dumps(expected, ensure_ascii=False), encoding="utf-8")
    runner = RegressionRunner(DocumentExtractor([_Stub(expected, fail_for={"bad"})]))
    summary = asyncio.run(runner.run(tmp_path)).to_dict()["summary"]
    assert summary["average_score"] == 50.0
    assert summary["scored_average_score"] == 100.0
    assert summary["coverage"] == 50.0


# ---------------------------------------------------------------------------
# B02 - regression evaluates the selected prompt
# ---------------------------------------------------------------------------

def test_b02_run_uses_the_selected_prompt_and_reports_it(client, monkeypatch, tmp_path):
    suite = tmp_path / "regression_suite"
    suite.mkdir()
    expected = _normalised(samples.SHL_VALID)
    (suite / "shl.pdf").write_bytes(PDF)
    (suite / "shl_expected.json").write_text(json.dumps(expected, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(get_settings(), "data_dir", str(tmp_path))

    seen = []
    monkeypatch.setattr(GeminiProvider, "is_configured", lambda self: True)

    def record(self, system_prompt, file_path, mime_type):
        seen.append(system_prompt)
        return ProviderResponse(raw_text=json.dumps(expected), model="gemini-test")
    monkeypatch.setattr(GeminiProvider, "extract", record)

    marker = f"MARKER-{uuid.uuid4().hex}"
    candidate = client.post("/api/prompts/", json={"name": "candidate", "content": marker}).json()
    report = client.post(f"/api/regression/run?prompt_id={candidate['id']}").json()["data"]
    assert marker in seen[-1]
    assert report["prompt"]["id"] == candidate["id"] and report["prompt"]["is_active"] is False
    assert report["providers"] == ["gemini/gemini-test"]

    active = client.get("/api/prompts/active").json()
    client.post("/api/regression/run")
    assert active["content"] in seen[-1] and marker not in seen[-1]

    assert client.post("/api/regression/run?prompt_id=999999").status_code == 404


def test_b02_no_real_provider_is_an_explicit_error(client):
    response = client.post("/api/regression/run")
    assert response.status_code == 409
    assert "GEMINI_API_KEY" in response.json()["detail"]


def test_b03_stale_claim_from_a_crashed_worker_can_be_retaken(client):
    from datetime import datetime, timedelta, timezone

    from models.processed_attachment import ProcessedAttachment
    from services.ingestion_pipeline import _claim_attachment

    fresh, stale = uuid.uuid4().hex, uuid.uuid4().hex

    async def scenario():
        async with async_session() as session:
            session.add(ProcessedAttachment(dedupe_key=fresh))
            session.add(ProcessedAttachment(dedupe_key=stale,
                                            claimed_at=datetime.now(timezone.utc) - timedelta(hours=1)))
            await session.commit()
        return (await _claim_attachment(fresh), await _claim_attachment(stale), await _claim_attachment(stale))

    in_progress, taken_over, second_taker = asyncio.run(scenario())
    assert in_progress == (False, None)
    assert taken_over == (True, None)
    assert second_taker == (False, None)


def test_review_edits_are_validated_and_deleting_allows_deliberate_reingest(client):
    message_id = f"<{uuid.uuid4().hex}@x>"
    first = client.post("/api/email/inbound/json", json=_json_email(message_id)).json()["results"][0]
    doc_id = first["ingestion_id"]

    edited = client.patch(f"/api/documents/{doc_id}", json={"extracted_data": {"total_amount": "NaN"}, "status": "VALID"})
    assert edited.status_code == 200
    stored = next(d for d in client.get("/api/documents?limit=500").json() if d["id"] == doc_id)
    assert stored["extracted_data"]["total_amount"] is None
    assert client.patch(f"/api/documents/{doc_id}", json={"status": "HACKED"}).status_code == 422

    assert client.delete(f"/api/documents/{doc_id}").status_code == 200
    again = client.post("/api/email/inbound/json", json=_json_email(message_id)).json()["results"][0]
    assert again["duplicate"] is False and again["status"] == "VALID"  # SQLite may reuse the id
