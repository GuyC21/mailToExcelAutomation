"""The Excel workbook that serves as GoldenCare's source of truth.

Writes are serialised with a lock and saved atomically (temp file + replace)
so a crash mid-write can never corrupt the accumulated workbook. If the file
is locked (typically: open in Excel on Windows) ``ExcelLockedError`` is raised
and the caller keeps the rows pending in the DB for the next sync.
"""
import logging
import os
import threading
from datetime import datetime
from typing import Iterable, List

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from repositories.excel_layout import (
    DOCUMENT_COLUMNS, DOCUMENTS_SHEET, LINE_COLUMNS, LINES_SHEET, STATUS_FILLS, STATUS_LABELS, Column,
)

logger = logging.getLogger(__name__)
_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_STATUS_COL = [c.header for c in DOCUMENT_COLUMNS].index("סטטוס") + 1


class ExcelLockedError(RuntimeError):
    """The workbook cannot be written right now (e.g. open in Excel)."""


class ExcelRepository:
    """Append-only access to the master workbook."""

    def __init__(self, file_path: str):
        self.file_path = file_path
        self._lock = threading.Lock()

    # ---- structure -------------------------------------------------------
    def ensure_workbook(self) -> None:
        """Creates the workbook, archiving any file with an outdated layout."""
        os.makedirs(os.path.dirname(self.file_path) or ".", exist_ok=True)
        with self._lock:
            if os.path.exists(self.file_path) and not self._has_current_layout():
                archived = self.file_path.replace(".xlsx", f".legacy-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
                os.replace(self.file_path, archived)
                logger.warning("Workbook layout outdated; archived to %s", archived)
            if not os.path.exists(self.file_path):
                self._save(self._new_workbook())

    def _has_current_layout(self) -> bool:
        workbook = load_workbook(self.file_path, read_only=True)
        try:
            if DOCUMENTS_SHEET not in workbook.sheetnames or LINES_SHEET not in workbook.sheetnames:
                return False
            header = next(workbook[DOCUMENTS_SHEET].iter_rows(max_row=1, values_only=True))
            return list(header) == [c.header for c in DOCUMENT_COLUMNS]
        finally:
            workbook.close()

    @staticmethod
    def _new_workbook() -> Workbook:
        workbook = Workbook()
        workbook.active.title = DOCUMENTS_SHEET
        workbook.create_sheet(LINES_SHEET)
        for title, columns in ((DOCUMENTS_SHEET, DOCUMENT_COLUMNS), (LINES_SHEET, LINE_COLUMNS)):
            sheet = workbook[title]
            sheet.sheet_view.rightToLeft = True
            sheet.freeze_panes = "A2"
            sheet.append([c.header for c in columns])
            for index, column in enumerate(columns, start=1):
                cell = sheet.cell(row=1, column=index)
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _HEADER_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                sheet.column_dimensions[cell.column_letter].width = column.width
            sheet.auto_filter.ref = sheet.dimensions
        return workbook

    # ---- writes ----------------------------------------------------------
    def append_records(self, records: List[dict]) -> List[int]:
        """Appends documents (+ their lines); skips ids already present.

        Returns:
            The ingestion ids now present in the workbook.
        """
        with self._lock:
            workbook = load_workbook(self.file_path)
            documents, lines = workbook[DOCUMENTS_SHEET], workbook[LINES_SHEET]
            existing = {row[0] for row in documents.iter_rows(min_row=2, max_col=1, values_only=True)}
            for record in records:
                if record["ingestion_id"] in existing:
                    continue  # Idempotent: a retry after a partial failure never duplicates rows.
                self._append_row(documents, DOCUMENT_COLUMNS, record)
                status_cell = documents.cell(row=documents.max_row, column=_STATUS_COL)
                status_cell.fill = PatternFill("solid", fgColor=STATUS_FILLS.get(record["status"], "FFFFFF"))
                for line in (record.get("data") or {}).get("line_items") or []:
                    self._append_row(lines, LINE_COLUMNS, {"record": record, "line": line})
            for sheet in (documents, lines):
                sheet.auto_filter.ref = sheet.dimensions
            self._save(workbook)
            return [r["ingestion_id"] for r in records]

    @staticmethod
    def _append_row(sheet, columns: Iterable[Column], source: dict) -> None:
        sheet.append([column.value(source) for column in columns])
        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=sheet.max_row, column=index)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column.number_format:
                cell.number_format = column.number_format

    def _save(self, workbook: Workbook) -> None:
        temp_path = f"{self.file_path}.tmp"
        try:
            workbook.save(temp_path)
            os.replace(temp_path, self.file_path)
        except OSError as error:  # PermissionError / EBUSY when Excel holds the file
            raise ExcelLockedError("קובץ האקסל נעול (כנראה פתוח באקסל). הנתונים נשמרו ויסונכרנו בהמשך.") from error
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    # ---- reads -----------------------------------------------------------
    def preview(self, limit: int = 10) -> dict:
        """Latest document rows (newest first) and per-status counts."""
        if not os.path.exists(self.file_path):
            return {"headers": [], "rows": [], "counts": {}, "total": 0, "updated_at": None}
        workbook = load_workbook(self.file_path, read_only=True)
        try:
            sheet = workbook[DOCUMENTS_SHEET]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        headers, body = list(rows[0]), rows[1:]
        status_index = _STATUS_COL - 1
        counts = {code: sum(1 for r in body if r[status_index] == label) for code, label in STATUS_LABELS.items()}
        latest = [[self._jsonable(v) for v in row] for row in reversed(body[-limit:])]
        updated = datetime.fromtimestamp(os.path.getmtime(self.file_path)).isoformat(timespec="seconds")
        return {"headers": headers, "rows": latest, "counts": counts, "total": len(body), "updated_at": updated}

    @staticmethod
    def _jsonable(value):
        return value.isoformat() if hasattr(value, "isoformat") else value
